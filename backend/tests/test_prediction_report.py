from datetime import datetime, timedelta, timezone
from decimal import Decimal

from openpyxl import load_workbook

from scripts.generate_prediction_report import (
    Forecast,
    MapRow,
    _stats,
    _team_stats,
    create_workbook,
    evaluate_maps,
)


NOW = datetime(2026, 10, 31, 12, tzinfo=timezone.utc)


def make_map(index=1, *, radiant_won=True, radiant="Radiant FC", dire="Dire FC", forecasts=None, tournament="Cup"):
    return MapRow(
        game_id=f"game-{index}", external_id=f"match-{index}", game_number=1,
        started_at=NOW, status="completed", radiant_won=radiant_won,
        radiant_team=radiant, dire_team=dire, tournament_id="cup-id",
        tournament=tournament, series_id=f"series-{index}", format="BO3",
        radiant_score=2 if radiant_won is True else None,
        dire_score=0 if radiant_won is True else None,
        forecasts=tuple(forecasts or ()),
    )


def forecast(radiant="0.7", dire="0.3", *, when=NOW - timedelta(minutes=10), role="default"):
    return Forecast(Decimal(radiant), Decimal(dire), "batru", role, when)


def test_correct_and_incorrect_predictions_with_each_side_favorite():
    evaluations = evaluate_maps([
        make_map(1, radiant_won=True, forecasts=[forecast()]),
        make_map(2, radiant_won=False, forecasts=[forecast()]),
        make_map(3, radiant_won=False, forecasts=[forecast("0.2", "0.8")]),
    ])
    assert [item.reason for item in evaluations] == [None, None, None]
    assert _stats(evaluations)["correct"] == 2
    assert _stats(evaluations)["incorrect"] == 1
    assert _team_stats(evaluations) == [
        ("Обычный", "Radiant FC", 2, 1, 1, 0.5),
        ("Обычный", "Dire FC", 1, 1, 0, 1.0),
    ]


def test_ambiguous_missing_result_and_missing_forecast_are_excluded():
    evaluations = evaluate_maps([
        make_map(1, forecasts=[forecast("0.5", "0.5")]),
        make_map(2, radiant_won=None, forecasts=[forecast()]),
        make_map(3),
        make_map(4, radiant="Radiant FC", dire=None, forecasts=[forecast()]),
    ])
    assert [item.reason for item in evaluations] == [
        "Неоднозначные вероятности", "Нет фактического результата", "Нет прогноза", "Невозможно сопоставить команды"
    ]
    assert _stats(evaluations)["valid"] == 0


def test_latest_forecast_is_used_even_after_map_start_or_without_map_time():
    older = forecast("0.6", "0.4", when=NOW - timedelta(minutes=20))
    newer = forecast("0.8", "0.2", when=NOW - timedelta(minutes=5))
    late = forecast("0.9", "0.1", when=NOW + timedelta(minutes=1))
    selected = evaluate_maps([make_map(forecasts=[older, newer, late])])
    assert selected[0].forecast == late
    assert selected[0].reason is None
    no_map_time = make_map(forecasts=[late])
    no_map_time = MapRow(**{**no_map_time.__dict__, "started_at": None})
    assert evaluate_maps([no_map_time])[0].forecast == late


def test_forecast_types_and_tournaments_remain_separate():
    items = evaluate_maps([
        make_map(1, forecasts=[forecast(), forecast("0.2", "0.8", role="abc")], tournament="Cup A"),
        make_map(2, forecasts=[forecast()], tournament="Cup B"),
    ])
    assert {item.kind for item in items} == {"Обычный", "С учётом ролей"}
    assert len([item for item in items if item.row.tournament == "Cup A"]) == 2
    assert len([item for item in items if item.row.tournament == "Cup B"]) == 1


def test_empty_workbook_and_tournament_sheets(tmp_path):
    output = tmp_path / "report.xlsx"
    create_workbook(evaluate_maps([make_map(forecasts=[forecast()])]), output)
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["Общая статистика", "Cup"]
    assert workbook["Общая статистика"]["H7"].number_format == "0.00%"
    ws = workbook["Общая статистика"]
    assert ws.tables
    assert ws.auto_filter.ref is None
    assert all(table.autoFilter is not None for table in ws.tables.values())
